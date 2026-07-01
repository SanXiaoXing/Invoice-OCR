from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from typing import List, Dict
import os
from datetime import datetime
from src.utils import amount_converter
from src.utils.logger import logger
from src.services.expense_calculator import calc_totals, build_preview_rows, SUBSIDY_PER_DAY


class ReportGenerator:
    def __init__(self):
        pass

    def generate_report(self, invoice_data_list: List[Dict], output_path: str,
                        days: int = 0) -> bool:
        try:
            logger.info(f"[ReportGenerator] 开始生成报销单: {output_path}")
            wb = Workbook()
            ws = wb.active
            ws.title = "差旅费用报销单"

            thin_border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )

            title_font = Font(name='微软雅黑', bold=True, size=18)
            header_font = Font(name='微软雅黑', bold=True, size=11)
            normal_font = Font(name='微软雅黑', size=10)
            bold_font = Font(name='微软雅黑', bold=True, size=10)

            row = 1

            # ── 标题行 ──
            ws.merge_cells(f'A{row}:I{row}')
            ws[f'A{row}'] = "差旅费用报销单"
            ws[f'A{row}'].font = title_font
            ws[f'A{row}'].alignment = Alignment(horizontal='center', vertical='center')
            row += 1

            # ── 报销日期 / 编号 ──
            ws[f'A{row}'] = f"报销日期：{datetime.now().strftime('%Y')}年{datetime.now().strftime('%m')}月{datetime.now().strftime('%d')}日"
            ws[f'A{row}'].font = normal_font
            ws[f'G{row}'] = "编号："
            ws[f'G{row}'].font = normal_font
            row += 1

            # ── 9 列表头 ──
            headers = ["出发地点", "到达地点", "交通金额", "飞机票", "住宿", "市内交通", "补助标准", "出差天数", "合计"]
            for col_idx, header in enumerate(headers, 1):
                cell = ws.cell(row=row, column=col_idx)
                cell.value = header
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.border = thin_border
            row += 1

            # ── 数据行（复用 expense_calculator） ──
            preview_rows = build_preview_rows(invoice_data_list, days)
            totals = calc_totals(invoice_data_list, days)

            for row_data in preview_rows:
                cells = row_data['cells']
                is_bold = row_data['bold']
                for col_idx, value in enumerate(cells, 1):
                    cell = ws.cell(row=row, column=col_idx)
                    if isinstance(value, float):
                        cell.value = round(value, 2)
                        cell.alignment = Alignment(horizontal='right', vertical='center')
                    else:
                        cell.value = str(value) if value else ''
                        cell.alignment = Alignment(horizontal='center', vertical='center')
                    cell.font = bold_font if is_bold else normal_font
                    cell.border = thin_border
                row += 1

            # ── 大写金额 ──
            row += 1
            ws[f'A{row}'] = "报销总额（大写）："
            ws[f'A{row}'].font = header_font
            ws.merge_cells(f'B{row}:F{row}')
            ws[f'B{row}'] = totals['chinese']
            ws[f'B{row}'].font = Font(name='微软雅黑', bold=True, size=12)
            ws[f'H{row}'] = f"¥{totals['total']:.2f}"
            ws[f'H{row}'].font = Font(name='微软雅黑', bold=True, size=12)
            ws[f'H{row}'].alignment = Alignment(horizontal='right')
            row += 1

            # ── 预借金额 / 退补金额 ──
            ws[f'A{row}'] = "预借金额："
            ws[f'A{row}'].font = normal_font
            ws.merge_cells(f'B{row}:C{row}')
            ws[f'B{row}'] = f"¥{totals['advance']:.2f}"
            ws[f'B{row}'].font = normal_font
            ws[f'B{row}'].alignment = Alignment(horizontal='left')

            ws[f'E{row}'] = "退/补金额："
            ws[f'E{row}'].font = normal_font
            ws.merge_cells(f'F{row}:G{row}')
            ws[f'F{row}'] = f"¥{totals['refund']:.2f}"
            ws[f'F{row}'].font = normal_font
            ws[f'F{row}'].alignment = Alignment(horizontal='left')
            row += 2

            # ── 附单据张数 ──
            train_count = sum(1 for d in invoice_data_list if d.get('type') == 'train')
            flight_count = sum(1 for d in invoice_data_list if d.get('type') == 'flight')
            hotel_count = sum(1 for d in invoice_data_list if d.get('type') == 'hotel')
            car_count = sum(1 for d in invoice_data_list if d.get('type') == 'car')
            invoice_count = sum(1 for d in invoice_data_list if d.get('type') == 'invoice')

            ws[f'A{row}'] = "附单据张数合计："
            ws[f'A{row}'].font = normal_font
            ws[f'D{row}'] = f"城际交通：{train_count + flight_count}"
            ws[f'D{row}'].font = normal_font
            ws[f'F{row}'] = f"其他：{hotel_count + car_count + invoice_count}"
            ws[f'F{row}'].font = normal_font
            row += 1

            # ── 列宽 ──
            col_widths = [12, 12, 12, 12, 12, 12, 12, 10, 14]
            for col_idx, width in enumerate(col_widths, 1):
                ws.column_dimensions[get_column_letter(col_idx)].width = width

            # ── 全局边框 & 对齐 ──
            for row_num in range(1, row + 1):
                for col_num in range(1, 10):
                    cell = ws.cell(row=row_num, column=col_num)
                    cell.border = thin_border
                    if not cell.alignment:
                        cell.alignment = Alignment(vertical='center', wrap_text=True)

            output_dir = os.path.dirname(output_path)
            if output_dir:
                os.makedirs(output_dir, exist_ok=True)
            wb.save(output_path)
            logger.info(f"[ReportGenerator] 报销单生成成功: {output_path}")
            return True
        except Exception as e:
            logger.error(f"[ReportGenerator] 生成报销单错误: {e}")
            import traceback
            traceback.print_exc()
            return False
